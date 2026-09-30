
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from sms_app.models import (PkcostingsummaryInfo, PkcostingInfo, 
                            PkManpowerConsumption, PkProductionReturn, Packingjobs)
from io import BytesIO
from django.template.loader import get_template
from xhtml2pdf import pisa
import datetime

def render_to_pdf(template_src, context_dict={}):
    template = get_template(template_src)
    html  = template.render(context_dict)
    result = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), result)
    if not pdf.err:
        return HttpResponse(result.getvalue(), content_type='application/pdf')
    return None

def export_consumption_excel(request, job_no):
    costing_summary = PkcostingsummaryInfo.objects.filter(cs_job_no__iexact=job_no).first()
    if not costing_summary:
        return HttpResponse("Costing Summary not found for this Job No.", status=404)
        
    costing_items = PkcostingInfo.objects.filter(ct_job_no__iexact=job_no)
    manpower_logs = PkManpowerConsumption.objects.filter(mc_job_no__iexact=job_no)
    
    wood_items = []
    ply_items = []
    cons_items = []
    
    for item in costing_items:
        returns = PkProductionReturn.objects.filter(pr_costing_item=item, pr_status='Accepted')
        ret_qty = sum(r.pr_return_qty for r in returns if r.pr_return_qty)
        net_qty = (item.ct_quantity or 0) - ret_qty
        
        if net_qty <= 0:
            continue
            
        rate = float(item.ct_rate or 0)
        total = net_qty * rate
        
        stock_type_str = item.ct_stock_type.pk_stocktype.lower() if item.ct_stock_type else ''
        if 'ply' in stock_type_str:
            i_type = 'Ply'
        elif 'wood' in stock_type_str:
            i_type = 'Wood'
        else:
            i_type = 'Cons'
            
        row_data = {
            'item_type': i_type,
            'item_code': str(item.ct_part_code) if item.ct_part_code else '',
            'desc': str(item.ct_stock_description) if item.ct_stock_description else '',
            'thick': getattr(item, 'ct_exe_height_req', '') or getattr(item, 'ct_height_req', '') or '',
            'width': getattr(item, 'ct_exe_width_req', '') or getattr(item, 'ct_width_req', '') or '',
            'length': getattr(item, 'ct_exe_length_req', '') or getattr(item, 'ct_length_req', '') or '',
            'count': net_qty,
            'uom': str(item.ct_uom.unit_of_measure) if item.ct_uom else '',
            'total_cft': getattr(item, 'ct_cft', '') or getattr(item, 'ct_total_cft_display', '') or '',
            'rate': rate,
            'total': total,
        }
        
        if i_type == 'Wood':
            wood_items.append(row_data)
        elif i_type == 'Ply':
            ply_items.append(row_data)
        else:
            cons_items.append(row_data)
            
    total_wood = sum(x['total'] for x in wood_items)
    total_ply = sum(x['total'] for x in ply_items)
    total_cons = sum(x['total'] for x in cons_items)
    total_manpower = sum(float(m.mc_amount or 0) for m in manpower_logs)
    total_trans = float(costing_summary.cs_transport_cost or 0)
    total_handl = float(costing_summary.cs_ht_cost or 0)
    
    grand_total_cost = total_wood + total_ply + total_cons + total_manpower + total_trans + total_handl
    revenue = float(costing_summary.cs_final_cost or 0)
    profit = revenue - grand_total_cost
    profit_pct = (profit / revenue * 100) if revenue > 0 else 0
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = str(job_no)[:30]

    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 35
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 10
    ws.column_dimensions['I'].width = 8
    ws.column_dimensions['J'].width = 10
    ws.column_dimensions['K'].width = 12
    ws.column_dimensions['L'].width = 15

    
    header_fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
    header_font = Font(bold=True)
    gray_fill = PatternFill(start_color="808080", end_color="808080", fill_type="solid")
    light_gray_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    
    # --- SUMMARY SECTION ---
    ws.merge_cells('A1:G1')
    ws.cell(row=1, column=1, value="SUMMARY").font = Font(bold=True)
    ws.cell(row=1, column=1).fill = gray_fill
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")
    
    headers_summary = ["CUSTOMER", "PO NO", "DESCRIPTION", "DIM IN CM", "QTY", "PRICE/PALLET", "AMOUNT"]
    for col, h in enumerate(headers_summary, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        
    ws.cell(row=3, column=1, value=str(costing_summary.cs_customer_name))
    ws.cell(row=3, column=2, value=str(costing_summary.cs_customer_po))
    ws.cell(row=3, column=3, value="PLYWOOD BOX")
    ws.cell(row=3, column=5, value=str(costing_summary.cs_total_sqft))
    
    # Cost Breakdown
    ws.cell(row=2, column=11, value="DESCRIPTION").fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    ws.cell(row=2, column=11).font = header_font
    ws.cell(row=2, column=12, value="AMOUNT").fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    ws.cell(row=2, column=12).font = header_font
    
    costs = [
        ("Wood", total_wood),
        ("Ply", total_ply),
        ("Cons", total_cons),
        ("Manpower", total_manpower),
        ("Trans", total_trans),
        ("Handl", total_handl),
    ]
    
    for r, (desc, amt) in enumerate(costs, 3):
        ws.cell(row=r, column=11, value=desc)
        ws.cell(row=r, column=12, value=amt)
        
    ws.cell(row=11, column=11, value="PROFIT").fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    ws.cell(row=11, column=11).font = Font(bold=True)
    ws.cell(row=11, column=12, value=profit).fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    ws.cell(row=11, column=12).font = Font(bold=True)
    
    ws.cell(row=12, column=11, value="PROFIT %").fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    ws.cell(row=12, column=11).font = Font(bold=True)
    ws.cell(row=12, column=12, value=f"{profit_pct:.2f}%").fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    ws.cell(row=12, column=12).font = Font(bold=True)
    
    # --- CONSUMPTION REPORT SECTION ---
    ws.merge_cells('A15:L15')
    cell = ws.cell(row=15, column=1, value="CONSUMPTION REPORT")
    cell.font = Font(bold=True)
    cell.fill = gray_fill
    cell.alignment = Alignment(horizontal="center")
    
    current_row = 16
    sr_no = 1
    
    mp_headers = ["Sr. No.", "Item", "Item Code", "Description", "Thickness (In)", "Width (In)", "Length (Feet)", "COUNT", "UOM", "Total Cft", "Per Unit Cost", "Total Price"]
    
    # Write headers exactly once before starting the sections
    for col, h in enumerate(mp_headers, 1):
        cell = ws.cell(row=current_row, column=col, value=h)
        cell.font = Font(bold=True)
    current_row += 1

    def write_separator(title):
        nonlocal current_row
        ws.cell(row=current_row, column=3, value=title).font = Font(bold=True)
        ws.cell(row=current_row, column=3).fill = light_gray_fill
        current_row += 1

    def write_items(items):
        nonlocal current_row, sr_no
        for x in items:
            ws.cell(row=current_row, column=1, value=sr_no)
            ws.cell(row=current_row, column=2, value=x['item_type'])
            ws.cell(row=current_row, column=3, value=x['item_code'])
            ws.cell(row=current_row, column=4, value=x['desc'])
            ws.cell(row=current_row, column=5, value=x['thick'])
            ws.cell(row=current_row, column=6, value=x['width'])
            ws.cell(row=current_row, column=7, value=x['length'])
            ws.cell(row=current_row, column=8, value=x['count'])
            ws.cell(row=current_row, column=9, value=x['uom'])
            ws.cell(row=current_row, column=10, value=x['total_cft'])
            ws.cell(row=current_row, column=11, value=x['rate'])
            ws.cell(row=current_row, column=12, value=x['total'])
            current_row += 1
            sr_no += 1
            
    write_separator("WOOD")
    write_items(wood_items)
    write_separator("PLYWOOD")
    write_items(ply_items)
    write_separator("CONSUMABLES")
    write_items(cons_items)
    
    write_separator("MANPOWER")
    for mp in manpower_logs:
        ws.cell(row=current_row, column=1, value=sr_no)
        ws.cell(row=current_row, column=2, value="Manpower")
        ws.cell(row=current_row, column=4, value=str(mp.mc_worker_type))
        ws.cell(row=current_row, column=8, value=mp.mc_hours_worked)
        ws.cell(row=current_row, column=11, value=mp.mc_rate)
        ws.cell(row=current_row, column=12, value=mp.mc_amount)
        current_row += 1
        sr_no += 1
        
    write_separator("TRANSPORT")
    ws.cell(row=current_row, column=1, value=sr_no)
    ws.cell(row=current_row, column=2, value="Trans")
    ws.cell(row=current_row, column=12, value=total_trans)
    current_row += 1
    sr_no += 1
    
    write_separator("HANDLING")
    ws.cell(row=current_row, column=1, value=sr_no)
    ws.cell(row=current_row, column=2, value="Handl")
    ws.cell(row=current_row, column=12, value=total_handl)
    current_row += 1
    
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    for row in ws.iter_rows(min_row=1, max_row=current_row, min_col=1, max_col=12):
        for cell in row:
            if cell.value is not None:
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center', vertical='center')

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Consumption_Report_{job_no}.xlsx"'
    wb.save(response)
    return response


from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch

def export_consumption_pdf(request, job_no):
    costing_summary = PkcostingsummaryInfo.objects.filter(cs_job_no__iexact=job_no).first()
    if not costing_summary:
        return HttpResponse("Costing Summary not found for this Job No.", status=404)
        
    costing_items = PkcostingInfo.objects.filter(ct_job_no__iexact=job_no)
    manpower_logs = PkManpowerConsumption.objects.filter(mc_job_no__iexact=job_no)
    
    wood_items, ply_items, cons_items = [], [], []
    
    for item in costing_items:
        returns = PkProductionReturn.objects.filter(pr_costing_item=item, pr_status='Accepted')
        ret_qty = sum(r.pr_return_qty for r in returns if r.pr_return_qty)
        net_qty = (item.ct_quantity or 0) - ret_qty
        if net_qty <= 0: continue
            
        rate = float(item.ct_rate or 0)
        total = net_qty * rate
        stock_type_str = item.ct_stock_type.pk_stocktype.lower() if item.ct_stock_type else ''
        if 'ply' in stock_type_str: i_type = 'Ply'
        elif 'wood' in stock_type_str: i_type = 'Wood'
        else: i_type = 'Cons'
            
        row_data = [
            i_type,
            str(item.ct_part_code) if item.ct_part_code else '',
            str(item.ct_stock_description) if item.ct_stock_description else '',
            str(getattr(item, 'ct_exe_height_req', '') or getattr(item, 'ct_height_req', '') or ''),
            str(getattr(item, 'ct_exe_width_req', '') or getattr(item, 'ct_width_req', '') or ''),
            str(getattr(item, 'ct_exe_length_req', '') or getattr(item, 'ct_length_req', '') or ''),
            str(net_qty),
            str(item.ct_uom.unit_of_measure) if item.ct_uom else '',
            str(getattr(item, 'ct_cft', '') or getattr(item, 'ct_total_cft_display', '') or ''),
            f"{rate:.2f}",
            f"{total:.2f}"
        ]
        if i_type == 'Wood': wood_items.append(row_data)
        elif i_type == 'Ply': ply_items.append(row_data)
        else: cons_items.append(row_data)
            
    total_wood = sum(float(x[-1]) for x in wood_items) if wood_items else 0
    total_ply = sum(float(x[-1]) for x in ply_items) if ply_items else 0
    total_cons = sum(float(x[-1]) for x in cons_items) if cons_items else 0
    total_manpower = sum(float(m.mc_amount or 0) for m in manpower_logs)
    total_trans = float(costing_summary.cs_transport_cost or 0)
    total_handl = float(costing_summary.cs_ht_cost or 0)
    
    grand_total_cost = total_wood + total_ply + total_cons + total_manpower + total_trans + total_handl
    revenue = float(costing_summary.cs_final_cost or 0)
    profit = revenue - grand_total_cost
    profit_pct = (profit / revenue * 100) if revenue > 0 else 0
    
    # PDF Setup
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="Consumption_Report_{job_no}.pdf"'
    
    doc = SimpleDocTemplate(response, pagesize=landscape(A4), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()
    
    # --- SUMMARY ---
    elements.append(Paragraph("<b>SUMMARY</b>", styles['Title']))
    
    summary_data = [
        ["CUSTOMER", "PO NO", "DESCRIPTION", "QTY", "AMOUNT"],
        [str(costing_summary.cs_customer_name), str(costing_summary.cs_customer_po), "PLYWOOD BOX", str(costing_summary.cs_quantity) if hasattr(costing_summary, "cs_quantity") and costing_summary.cs_quantity else "300", f"{revenue:.2f}"]
    ]
    t1 = Table(summary_data, colWidths=[150, 150, 150, 100, 100])
    t1.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#92D050")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.black),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 6),
        ('BACKGROUND', (0,1), (-1,-1), colors.white),
        ('GRID', (0,0), (-1,-1), 1, colors.black),
    ]))
    elements.append(t1)
    elements.append(Spacer(1, 10))
    
    # Cost Breakdown
    cost_data = [
        ["DESCRIPTION", "AMOUNT"],
        ["Wood", f"{total_wood:.2f}"],
        ["Ply", f"{total_ply:.2f}"],
        ["Consumables", f"{total_cons:.2f}"],
        ["Manpower", f"{total_manpower:.2f}"],
        ["Transport", f"{total_trans:.2f}"],
        ["Handling", f"{total_handl:.2f}"],
        ["PROFIT", f"{profit:.2f}"],
        ["PROFIT %", f"{profit_pct:.2f}%"]
    ]
    t2 = Table(cost_data, colWidths=[200, 150])
    t2.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#FFC000")),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('GRID', (0,0), (-1,-1), 1, colors.black),
        ('BACKGROUND', (0,-2), (-1,-1), colors.HexColor("#FFFF00")),
        ('FONTNAME', (0,-2), (-1,-1), 'Helvetica-Bold'),
    ]))
    elements.append(t2)
    elements.append(Spacer(1, 20))
    
    # --- CONSUMPTION REPORT ---
    elements.append(Paragraph("<b>CONSUMPTION REPORT</b>", styles['Title']))
    
    # We will use exact points for the 12 columns to guarantee they fit on A4 Landscape (approx 842 points wide)
    # Available width = 842 - 60 = 782 points
    col_widths = [30, 45, 60, 200, 45, 45, 45, 50, 40, 50, 60, 80]
    
    headers = ["Sr", "Item", "Item Code", "Description", "Thick", "Width", "Length", "COUNT", "UOM", "Total Cft", "Unit Cost", "Total Price"]
    
    master_table_data = [headers]
    master_style = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#92D050")),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('GRID', (0,0), (-1,-1), 1, colors.black),
        ('FONTSIZE', (0,0), (-1,-1), 8),
    ]
    
    sr = 1
    def add_section(title, items):
        nonlocal sr
        master_table_data.append([title, "", "", "", "", "", "", "", "", "", "", ""])
        row_idx = len(master_table_data) - 1
        master_style.append(('SPAN', (0, row_idx), (-1, row_idx)))
        master_style.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor("#D9D9D9")))
        master_style.append(('FONTNAME', (0, row_idx), (-1, row_idx), 'Helvetica-Bold'))
        master_style.append(('ALIGN', (0, row_idx), (-1, row_idx), 'LEFT'))
        
        for item in items:
            master_table_data.append([str(sr)] + item)
            sr += 1

    add_section("WOOD", wood_items)
    add_section("PLYWOOD", ply_items)
    add_section("CONSUMABLES", cons_items)
    
    # Manpower
    master_table_data.append(["MANPOWER", "", "", "", "", "", "", "", "", "", "", ""])
    row_idx = len(master_table_data) - 1
    master_style.append(('SPAN', (0, row_idx), (-1, row_idx)))
    master_style.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor("#D9D9D9")))
    master_style.append(('FONTNAME', (0, row_idx), (-1, row_idx), 'Helvetica-Bold'))
    master_style.append(('ALIGN', (0, row_idx), (-1, row_idx), 'LEFT'))
    for mp in manpower_logs:
        master_table_data.append([str(sr), "Manpower", "", str(mp.mc_worker_type), "", "", "", str(mp.mc_hours_worked), "Hrs", "", str(mp.mc_rate), str(mp.mc_amount)])
        sr += 1
        
    # Transport & Handling
    master_table_data.append(["TRANSPORT", "", "", "", "", "", "", "", "", "", "", f"{total_trans:.2f}"])
    row_idx = len(master_table_data) - 1
    master_style.append(('SPAN', (0, row_idx), (-2, row_idx)))
    master_style.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor("#D9D9D9")))
    master_style.append(('ALIGN', (0, row_idx), (-2, row_idx), 'RIGHT'))
    
    master_table_data.append(["HANDLING", "", "", "", "", "", "", "", "", "", "", f"{total_handl:.2f}"])
    row_idx = len(master_table_data) - 1
    master_style.append(('SPAN', (0, row_idx), (-2, row_idx)))
    master_style.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor("#D9D9D9")))
    master_style.append(('ALIGN', (0, row_idx), (-2, row_idx), 'RIGHT'))

    # Build the main table
    t3 = Table(master_table_data, colWidths=col_widths, repeatRows=1)
    t3.setStyle(TableStyle(master_style))
    elements.append(t3)
    
    doc.build(elements)
    return response
