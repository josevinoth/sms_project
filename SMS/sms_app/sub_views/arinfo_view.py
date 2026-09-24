from django.contrib.auth.decorators import login_required
from django.views.decorators.csrf import csrf_exempt
from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse, HttpResponse
from django.db import connection
from django.urls import reverse
from django.utils import timezone
from datetime import datetime, date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..forms import ArinfoaddForm
from ..models import (
    BilingInfo, Ar_comments_Info, User_extInfo, Ar_Info,
    TransInvoiceInfo, CustomerInfo, Emailmaster, Location_info,
    Business_Sol_info, Bvmproduct
)


def _get_submission_map():
    sub_map = {}
    with connection.cursor() as cursor:
        cursor.execute("SELECT ti_inv_no, ti_submission_date, ti_submitted_to FROM sms_app_transinvoiceinfo WHERE ti_submission_date IS NOT NULL")
        for inv_no, sub_dt, sub_to in cursor.fetchall():
            if inv_no and inv_no not in sub_map:
                sub_map[inv_no] = {'date': sub_dt, 'to': sub_to or ''}
    return sub_map


def _calculate_aging(sub_date, credit_days, current_date=None):
    """
    Calculates aging status and color coding for submitted invoices:
    - days = (current_date - sub_date).days
    - credit_days: customer credit days from CustomerInfo (fallback to 30)
    - If days <= credit_days: Not Due (Green: #059669 / #ecfdf5 / #10b981)
    - If credit_days < days <= 2 * credit_days: Due (Amber/Orange: #d97706 / #fffbeb / #f59e0b)
    - If days > 2 * credit_days: Overdue (Red: #dc2626 / #fef2f2 / #ef4444)
    """
    if not sub_date:
        return None

    if current_date is None:
        current_date = timezone.localtime(timezone.now()).date()

    if isinstance(sub_date, str):
        try:
            sub_date = datetime.strptime(sub_date.split('T')[0], '%Y-%m-%d').date()
        except Exception:
            return None
    elif hasattr(sub_date, 'date') and not isinstance(sub_date, date):
        sub_date = sub_date.date()

    try:
        cdays = int(credit_days) if credit_days is not None and int(credit_days) > 0 else 30
    except (ValueError, TypeError):
        cdays = 30

    days_diff = (current_date - sub_date).days
    if days_diff < 0:
        days_diff = 0

    if days_diff <= cdays:
        status_label = 'Not Due'
        status_code = 'not_due'
        color = '#059669'
        bg_color = '#ecfdf5'
        border_color = '#10b981'
        badge_class = 'badge-aging-not-due'
    elif days_diff <= 2 * cdays:
        status_label = 'Due'
        status_code = 'due'
        color = '#d97706'
        bg_color = '#fffbeb'
        border_color = '#f59e0b'
        badge_class = 'badge-aging-due'
    else:
        status_label = 'Overdue'
        status_code = 'overdue'
        color = '#dc2626'
        bg_color = '#fef2f2'
        border_color = '#ef4444'
        badge_class = 'badge-aging-overdue'

    return {
        'days': days_diff,
        'credit_days': cdays,
        'status_label': status_label,
        'status_code': status_code,
        'color': color,
        'bg_color': bg_color,
        'border_color': border_color,
        'badge_class': badge_class,
    }


def _get_customer_to_cc_recipients(customer):
    """
    Fetch separated To and Cc auto-recipients for a customer from Emailmaster and CustomerInfo.
    """
    if not customer:
        return "", ""
    to_list = []
    cc_list = []
    em_qs = Emailmaster.objects.filter(em_Customer_name=customer)
    for em in em_qs:
        if em.em_to_names:
            to_list.extend([e.strip() for e in em.em_to_names.split(',') if e.strip()])
        if em.em_cc_names:
            cc_list.extend([e.strip() for e in em.em_cc_names.split(',') if e.strip()])
    if getattr(customer, 'cu_email', None):
        to_list.extend([e.strip() for e in customer.cu_email.split(',') if e.strip()])

    def dedupe(lst):
        seen = set()
        cleaned = []
        for r in lst:
            r_clean = r.strip()
            if r_clean and r_clean.lower() not in seen:
                seen.add(r_clean.lower())
                cleaned.append(r_clean)
        return cleaned

    return ", ".join(dedupe(to_list)), ", ".join(dedupe(cc_list))


def _build_combined_pdf_for_invoices(inv_no_list):
    """
    Builds a single combined PDF (merging master uploaded PDF, Ann-1, Ann-2, Ann-3)
    for all given invoice numbers.
    """
    from pypdf import PdfWriter, PdfReader
    from .trans_invoice_view import _render_ann1_pdf_bytes, _render_cnote_pdf_bytes, _sync_trans_invoice_pdf
    import io

    combined_writer = PdfWriter()
    total_pages = 0

    for inv_no in inv_no_list:
        master_inv = TransInvoiceInfo.objects.filter(ti_inv_no=inv_no, is_woh=False).first()
        if not master_inv:
            master_inv = TransInvoiceInfo.objects.filter(ti_inv_no=inv_no).first()
        if not master_inv:
            continue

        try:
            _sync_trans_invoice_pdf(master_inv.ti_inv_no, master_inv.ti_customer_id)
            master_inv.refresh_from_db()
        except Exception:
            pass

        # 1. Primary uploaded invoice PDF
        if master_inv.ti_invoice_pdf:
            try:
                master_inv.ti_invoice_pdf.open('rb')
                reader = PdfReader(master_inv.ti_invoice_pdf)
                for page in reader.pages:
                    combined_writer.add_page(page)
                    total_pages += 1
            except Exception as e:
                print(f"Error reading ti_invoice_pdf for {inv_no}: {e}")
            finally:
                try:
                    master_inv.ti_invoice_pdf.close()
                except Exception:
                    pass

        # 2. Ann-1 (Billing Statement)
        try:
            ann1_bytes = _render_ann1_pdf_bytes(inv_no)
            if ann1_bytes:
                reader = PdfReader(io.BytesIO(ann1_bytes))
                for page in reader.pages:
                    combined_writer.add_page(page)
                    total_pages += 1
        except Exception as e:
            print(f"Error rendering Ann-1 for {inv_no}: {e}")

        # 3. Ann-2 (Combined CNote)
        try:
            ann2_bytes = _render_cnote_pdf_bytes(master_inv)
            if ann2_bytes:
                reader = PdfReader(io.BytesIO(ann2_bytes))
                for page in reader.pages:
                    combined_writer.add_page(page)
                    total_pages += 1
        except Exception as e:
            print(f"Error rendering Ann-2 for {inv_no}: {e}")

        # 4. Ann-3 (Merged invoice documents)
        if master_inv.ti_merged_pdf:
            try:
                master_inv.ti_merged_pdf.open('rb')
                reader = PdfReader(master_inv.ti_merged_pdf)
                for page in reader.pages:
                    combined_writer.add_page(page)
                    total_pages += 1
            except Exception as e:
                print(f"Error reading ti_merged_pdf for {inv_no}: {e}")
            finally:
                try:
                    master_inv.ti_merged_pdf.close()
                except Exception:
                    pass

    if total_pages > 0:
        out = io.BytesIO()
        combined_writer.write(out)
        out.seek(0)
        return out.getvalue()
    return None


@login_required(login_url='login_page')
def ar_list(request):
    first_name = request.session.get('first_name')
    user_id = request.session.get('ses_userID')
    role = None
    try:
        role = User_extInfo.objects.get(user=user_id).emp_role
    except Exception:
        pass

    today = timezone.localtime(timezone.now()).date()

    # Pre-build planning date mapping from saved invoice line items (is_woh=True)
    woh_planning_dates = (
        TransInvoiceInfo.objects
        .filter(is_woh=True)
        .values_list(
            'ti_inv_no',
            'ti_trip__tr_enquirynumber__en_pickupdatetime',
            'ti_consignment__co_enquirynumber__en_pickupdatetime'
        )
    )
    plan_map = {}
    for inv_no, trip_dt, cons_dt in woh_planning_dates:
        if inv_no:
            dt = trip_dt or cons_dt
            if dt:
                d = dt.date() if hasattr(dt, 'date') else dt
                if inv_no not in plan_map or d > plan_map[inv_no]:
                    plan_map[inv_no] = d

    # Pre-build department mapping from WOH line items (ti_department populated on WOH records)
    woh_dept_qs = (
        TransInvoiceInfo.objects
        .filter(is_woh=True)
        .exclude(ti_department__isnull=True)
        .exclude(ti_department='')
        .values_list('ti_inv_no', 'ti_department')
    )
    dept_map = {}
    for inv_no, dept in woh_dept_qs:
        if inv_no and inv_no not in dept_map:
            dept_map[inv_no] = dept

    # Pre-build submission tracking map via raw SQL (without model changes)
    sub_map = _get_submission_map()

    # Fetch Transport Master Invoices (auto-fetched submitted invoices)
    trans_qs = (
        TransInvoiceInfo.objects
        .filter(is_woh=False)
        .select_related(
            'ti_customer', 'ti_customer__cu_business_sol',
            'ti_trip', 'ti_trip__tr_enquirynumber',
            'ti_consignment', 'ti_consignment__co_enquirynumber',
            'ti_goods'
        )
        .order_by('-ti_inv_date', '-id')
    )

    unified_list = []
    total_val = 0.0
    trans_count = 0

    # Process TransInvoiceInfo (without modifying TransInvoiceInfo model)
    for ti in trans_qs:
        trans_count += 1
        amt = float(ti.ti_total or 0.0)
        total_val += amt

        # Strictly take ONLY the Planning Date (en_pickupdatetime):
        # 1. From saved line items (plan_map)
        # 2. From direct trip/consignment on master record
        # No fallback to invoice date
        op_date = plan_map.get(ti.ti_inv_no)
        if not op_date:
            try:
                if ti.ti_trip and ti.ti_trip.tr_enquirynumber and ti.ti_trip.tr_enquirynumber.en_pickupdatetime:
                    pdt = ti.ti_trip.tr_enquirynumber.en_pickupdatetime
                    op_date = pdt.date() if hasattr(pdt, 'date') else pdt
                elif ti.ti_consignment and ti.ti_consignment.co_enquirynumber and ti.ti_consignment.co_enquirynumber.en_pickupdatetime:
                    pdt = ti.ti_consignment.co_enquirynumber.en_pickupdatetime
                    op_date = pdt.date() if hasattr(pdt, 'date') else pdt
            except Exception:
                op_date = None

        comp_name = 'BVM Trans Solutions Pvt Ltd'
        if ti.ti_customer and ti.ti_customer.cu_business_sol:
            comp_name = str(ti.ti_customer.cu_business_sol)

        # Customer Name: use cu_name from CustomerInfo (e.g. CRANE(T)MAA)
        c_name = (
            (getattr(ti.ti_customer, 'cu_name', None) if ti.ti_customer else '') or
            ti.ti_customer_short_name or
            (str(ti.ti_customer) if ti.ti_customer else '')
        )

        # Due calculations
        due_op = (today - op_date).days if op_date else ''
        due_sub = ''  # Kept empty per requirement

        # Customer department: fetch from WOH line items (dept_map), fallback to master field
        c_dept = dept_map.get(ti.ti_inv_no) or ti.ti_department or ''

        sub_info = sub_map.get(ti.ti_inv_no)
        sub_date = sub_info['date'] if sub_info else None
        sub_to = sub_info['to'] if sub_info else ''

        is_submitted = bool(sub_date)
        if is_submitted:
            status_label = 'Submitted'
        else:
            status_label = 'Not Submitted'

        has_pdf = bool(ti.ti_invoice_pdf and ti.ti_invoice_pdf.name)

        aging = None
        if is_submitted and sub_date:
            credit_days = ti.ti_customer.cu_creditdays if ti.ti_customer else None
            aging = _calculate_aging(sub_date, credit_days, today)

        unified_list.append({
            'unique_id': f"trans_{ti.id}",
            'source': 'trans',
            'db_id': ti.id,
            'company': comp_name,
            'product': 'Transport',
            'branch': ti.ti_branch or '',
            'operation_date': op_date,
            'invoice_num': ti.ti_inv_no or '',
            'invoice_date': ti.ti_inv_date,
            'customer_name': c_name,
            'customer_id': ti.ti_customer_id,
            'customer_dept': c_dept,
            'service_value': amt,
            'cgst': 0.0,
            'sgst': 0.0,
            'igst': 0.0,
            'amount': amt + 0.0 + 0.0 + 0.0,
            'due_from_operation_date': due_op,
            'due_from_submission_date': '',
            'submission_date': sub_date,
            'submitted_to': sub_to,
            'payment_received_date': None,
            'payment_received_amount': 0.0,
            'tds': 0.0,
            'sales_person': '',
            'status': status_label,
            'has_invoice_pdf': has_pdf,
            'aging': aging,
        })

    # Sort unified list: invoices with most recent date first
    unified_list.sort(key=lambda x: str(x['invoice_date'] or x['operation_date'] or ''), reverse=True)

    context = {
        'ar_list': unified_list,
        'first_name': first_name,
        'role': role,
        'metrics': {
            'total_invoices': len(unified_list),
            'total_value': round(total_val, 2),
            'trans_count': trans_count,
        }
    }
    return render(request, "asset_mgt_app/ar_list.html", context)


def _get_dept_map():
    woh_dept_qs = (
        TransInvoiceInfo.objects
        .filter(is_woh=True)
        .exclude(ti_department__isnull=True)
        .exclude(ti_department='')
        .values_list('ti_inv_no', 'ti_department')
    )
    dept_map = {}
    for inv_no, dept in woh_dept_qs:
        if inv_no and inv_no not in dept_map:
            dept_map[inv_no] = dept
    return dept_map


@login_required(login_url='login_page')
def fetch_invoice_submission_details(request):
    """
    Given a list of unique_ids (e.g. ['trans_16', 'ar_4']),
    fetches To/Cc auto-recipient list from Emailmaster & CustomerInfo,
    determines if all are already submitted (Outstanding Followup) or not (Invoice Submission),
    provides default subject and mail content, plus line items summary.
    """
    raw_ids = request.GET.get('ids', '') or request.POST.get('ids', '')
    if not raw_ids:
        return JsonResponse({'success': False, 'msg': 'No invoices selected.'})

    id_list = [x.strip() for x in raw_ids.split(',') if x.strip()]
    if not id_list:
        return JsonResponse({'success': False, 'msg': 'No valid IDs provided.'})

    invoices_data = []
    customer_ids = set()
    existing_dates = []
    existing_tos = []

    sub_map = _get_submission_map()
    dept_map = _get_dept_map()

    for item_id in id_list:
        if item_id.startswith('trans_'):
            try:
                ti_id = int(item_id.replace('trans_', ''))
                ti = TransInvoiceInfo.objects.select_related('ti_customer').get(id=ti_id)
                sub_info = sub_map.get(ti.ti_inv_no)
                s_date = str(sub_info['date']) if sub_info and sub_info['date'] else ''
                s_to = str(sub_info['to']) if sub_info and sub_info['to'] else ''
                c_dept = dept_map.get(ti.ti_inv_no) or ti.ti_department or '-'
                inv_date_str = ti.ti_inv_date.strftime('%d/%m/%Y') if ti.ti_inv_date else '-'

                cust_display = (
                    (getattr(ti.ti_customer, 'cu_name', None) if ti.ti_customer else '') or
                    ti.ti_customer_short_name or
                    (str(ti.ti_customer) if ti.ti_customer else '-')
                )

                has_pdf = bool(ti.ti_invoice_pdf and ti.ti_invoice_pdf.name)
                credit_days = ti.ti_customer.cu_creditdays if ti.ti_customer else None
                aging = _calculate_aging(sub_info['date'] if sub_info else None, credit_days)

                invoices_data.append({
                    'id': item_id,
                    'inv_no': ti.ti_inv_no or '-',
                    'inv_date': inv_date_str,
                    'customer': str(cust_display),
                    'department': c_dept,
                    'amount': float(ti.ti_total or 0.0),
                    'submission_date': s_date,
                    'submitted_to': s_to,
                    'has_pdf': has_pdf,
                    'aging': aging,
                })
                if ti.ti_customer:
                    customer_ids.add(ti.ti_customer)
                if s_date:
                    existing_dates.append(s_date)
                if s_to:
                    existing_tos.append(s_to)
            except Exception as e:
                print('Error fetching trans invoice detail:', e)
        elif item_id.startswith('ar_'):
            try:
                ar_id = int(item_id.replace('ar_', ''))
                ar = Ar_Info.objects.select_related('ar_customer_name', 'ar_invoice_num', 'ar_product').get(id=ar_id)
                inv_date_str = ar.ar_invoice_date.strftime('%d/%m/%Y') if getattr(ar, 'ar_invoice_date', None) else '-'
                s_date = str(ar.ar_submission_date) if ar.ar_submission_date else ''
                s_to = str(ar.ar_invoice_sent_to) if getattr(ar, 'ar_invoice_sent_to', None) else ''
                aging = _calculate_aging(ar.ar_submission_date, 30)
                invoices_data.append({
                    'id': item_id,
                    'inv_no': str(ar.ar_invoice_num or '-'),
                    'inv_date': inv_date_str,
                    'customer': str(ar.ar_customer_name or '-'),
                    'department': str(getattr(ar, 'ar_department', None) or '-'),
                    'amount': float(ar.ar_amount or 0.0),
                    'submission_date': s_date,
                    'submitted_to': s_to,
                    'has_pdf': True,
                    'aging': aging,
                })
                if ar.ar_customer_name:
                    customer_ids.add(ar.ar_customer_name)
                if s_date:
                    existing_dates.append(s_date)
                if s_to:
                    existing_tos.append(s_to)
            except Exception as e:
                print('Error fetching AR detail:', e)

    # Build To and Cc lists
    all_to = []
    all_cc = []
    for cust in customer_ids:
        to_part, cc_part = _get_customer_to_cc_recipients(cust)
        if to_part:
            all_to.extend([e.strip() for e in to_part.split(',') if e.strip()])
        if cc_part:
            all_cc.extend([e.strip() for e in cc_part.split(',') if e.strip()])

    def dedupe(lst):
        seen = set()
        res = []
        for item in lst:
            clean = item.strip()
            if clean and clean.lower() not in seen:
                seen.add(clean.lower())
                res.append(clean)
        return res

    final_to = dedupe(all_to)
    final_cc = dedupe(all_cc)

    # Determine mail type: if ALL selected invoices are already submitted -> Outstanding Followup
    is_all_submitted = (len(existing_dates) > 0 and len(existing_dates) == len(invoices_data))

    if is_all_submitted:
        mail_type = 'outstanding'
        default_subject = "Outstanding Followup"
        default_body = (
            "Dear Team,\n\n"
            "Please find the below Outstanding invoices \n\n"
            "Request you to clear the Due & Overdue payments\n\n"
            "The Invoices are highlighted as below"
        )
        default_submitted_to = existing_tos[0] if existing_tos else ''
    else:
        mail_type = 'submission'
        default_subject = "Invoice Submission"
        default_body = (
            "Dear Team,\n\n"
            "Please find the attached invoice details for your further process\n\n"
            "Request you to please acknowledge the receipt of the below invoices"
        )
        default_submitted_to = ''  # MUST STAY EMPTY BY DEFAULT FOR NEW SUBMISSION

        # Check if any selected invoice does NOT have an invoice PDF uploaded
        missing_pdf_invoices = [inv['inv_no'] for inv in invoices_data if not inv.get('has_pdf')]
        if missing_pdf_invoices:
            inv_str = ", ".join(missing_pdf_invoices)
            return JsonResponse({
                'success': False,
                'msg': f"The selected invoice ({inv_str}) doesn't have an invoice file."
            })

    # Default submission date: common date if all match, else today
    default_sub_date = timezone.localtime(timezone.now()).date().strftime('%Y-%m-%d')
    if existing_dates and len(set(existing_dates)) == 1:
        default_sub_date = existing_dates[0]

    return JsonResponse({
        'success': True,
        'count': len(invoices_data),
        'mail_type': mail_type,
        'subject': default_subject,
        'body': default_body,
        'submission_date': default_sub_date,
        'submitted_to': default_submitted_to,
        'to_emails': ", ".join(final_to),
        'cc_emails': ", ".join(final_cc),
        'invoices': invoices_data,
        'total_amount': sum(item['amount'] for item in invoices_data),
    })


@csrf_exempt
@login_required(login_url='login_page')
def ar_submit_selected_invoices(request):
    """
    Submits selected invoices:
    - Saves submitted_date and submitted_to to DB (for trans invoice and line items)
    - Sends email with custom Subject, separated To and Cc, Mail Content, invoice table,
      and attaches combined PDF for the selected invoices.
    """
    if request.method != "POST":
        return JsonResponse({'success': False, 'msg': 'Invalid request method.'})

    raw_ids = request.POST.get('selected_ids', '')
    submitted_date_str = request.POST.get('submitted_date', '').strip()
    submitted_to = request.POST.get('submitted_to', '').strip()
    to_str = request.POST.get('to_emails', '').strip() or request.POST.get('recipient', '').strip()
    cc_str = request.POST.get('cc_emails', '').strip()
    mail_subject = request.POST.get('mail_subject', '').strip()
    mail_body = request.POST.get('mail_body', '').strip()
    mail_type = request.POST.get('mail_type', '').strip()
    id_list = [x.strip() for x in raw_ids.split(',') if x.strip()]
    if not id_list:
        return JsonResponse({'success': False, 'msg': 'No invoices selected.'})

    # REQUIREMENT: For invoice submission (not outstanding), Submitted To MUST NOT be empty
    if mail_type != 'outstanding' and not submitted_to:
        return JsonResponse({
            'success': False,
            'msg': 'Submitted To is a required field. Please fill in Submitted To before submitting.'
        })

    # REQUIREMENT: Check whether invoice file is uploaded in trans invoice list
    if mail_type != 'outstanding':
        missing_pdf_invoices = []
        for item_id in id_list:
            if item_id.startswith('trans_'):
                try:
                    ti_id = int(item_id.replace('trans_', ''))
                    ti_chk = TransInvoiceInfo.objects.filter(id=ti_id).first()
                    if ti_chk and not (ti_chk.ti_invoice_pdf and ti_chk.ti_invoice_pdf.name):
                        missing_pdf_invoices.append(ti_chk.ti_inv_no or item_id)
                except Exception:
                    pass
        if missing_pdf_invoices:
            inv_str = ", ".join(missing_pdf_invoices)
            return JsonResponse({
                'success': False,
                'msg': f"The selected invoice ({inv_str}) doesn't have an invoice file."
            })

    to_recipients = [e.strip() for e in to_str.split(',') if e.strip()]
    cc_recipients = [e.strip() for e in cc_str.split(',') if e.strip()]
    if not to_recipients:
        return JsonResponse({
            'success': False,
            'msg': 'Please enter at least one recipient email in the To field.'
        })

    parsed_sub_date = None
    if submitted_date_str:
        try:
            parsed_sub_date = datetime.strptime(submitted_date_str, '%Y-%m-%d').date()
        except Exception:
            parsed_sub_date = timezone.localtime(timezone.now()).date()
    else:
        parsed_sub_date = timezone.localtime(timezone.now()).date()

    dept_map = _get_dept_map()
    sub_map = _get_submission_map()
    updated_count = 0
    preview_rows = []
    trans_inv_nos = []
    total_amt = 0.0

    for item_id in id_list:
        if item_id.startswith('trans_'):
            try:
                ti_id = int(item_id.replace('trans_', ''))
                ti = TransInvoiceInfo.objects.select_related('ti_customer').get(id=ti_id)
                inv_no = ti.ti_inv_no
                if inv_no and inv_no not in trans_inv_nos:
                    trans_inv_nos.append(inv_no)

                # Update master record and ALL line items via raw SQL if submission
                if mail_type != 'outstanding' or submitted_to:
                    with connection.cursor() as cursor:
                        cursor.execute(
                            "UPDATE sms_app_transinvoiceinfo SET ti_submission_date = %s, ti_submitted_to = %s WHERE ti_inv_no = %s",
                            [parsed_sub_date, submitted_to, inv_no]
                        )
                updated_count += 1
                c_dept = dept_map.get(inv_no) or ti.ti_department or '-'
                c_name = (
                    (getattr(ti.ti_customer, 'cu_name', None) if ti.ti_customer else '') or
                    ti.ti_customer_short_name or
                    (str(ti.ti_customer) if ti.ti_customer else '-')
                )
                amt = float(ti.ti_total or 0.0)
                total_amt += amt
                sub_info = sub_map.get(inv_no)
                s_date = (sub_info['date'] if sub_info else None) or parsed_sub_date
                c_creditdays = ti.ti_customer.cu_creditdays if ti.ti_customer else None
                aging = _calculate_aging(s_date, c_creditdays, timezone.localtime(timezone.now()).date())

                preview_rows.append({
                    'inv_no': inv_no or '-',
                    'inv_date': ti.ti_inv_date.strftime('%d/%m/%Y') if ti.ti_inv_date else '-',
                    'customer': c_name,
                    'department': c_dept,
                    'amount': amt,
                    'aging': aging,
                })
            except Exception as e:
                print('Error submitting trans invoice:', e)

        elif item_id.startswith('ar_'):
            try:
                ar_id = int(item_id.replace('ar_', ''))
                ar = Ar_Info.objects.select_related('ar_customer_name').get(id=ar_id)
                if mail_type != 'outstanding' or submitted_to:
                    ar.ar_submission_date = parsed_sub_date
                    ar.ar_invoice_sent_to = submitted_to
                    if ar.ar_operation_date:
                        ar.ar_due_from_operation_date = (parsed_sub_date - ar.ar_operation_date).days
                    ar.save()
                updated_count += 1
                amt = float(ar.ar_amount or 0.0)
                total_amt += amt
                inv_date_str = ar.ar_invoice_date.strftime('%d/%m/%Y') if getattr(ar, 'ar_invoice_date', None) else '-'
                aging = _calculate_aging(ar.ar_submission_date or parsed_sub_date, 30, timezone.localtime(timezone.now()).date())
                preview_rows.append({
                    'inv_no': str(ar.ar_invoice_num or '-'),
                    'inv_date': inv_date_str,
                    'customer': str(ar.ar_customer_name or '-'),
                    'department': str(getattr(ar, 'ar_department', None) or '-'),
                    'amount': amt,
                    'aging': aging,
                })
            except Exception as e:
                print('Error submitting AR invoice:', e)

    # Subject and Mail Content
    subject = mail_subject or ("Outstanding Followup" if mail_type == 'outstanding' else "Invoice Submission")

    if not mail_body:
        if mail_type == 'outstanding':
            mail_body = (
                "Dear Team,\n\n"
                "Please find the below Outstanding invoices \n\n"
                "Request you to clear the Due & Overdue payments\n\n"
                "The Invoices are highlighted as below"
            )
        else:
            mail_body = (
                "Dear Team,\n\n"
                "Please find the attached invoice details for your further process\n\n"
                "Request you to please acknowledge the receipt of the below invoices"
            )

    paragraphs = [p.replace('\n', '<br>') for p in mail_body.split('\n\n') if p.strip()]
    formatted_body_html = "".join([f"<p style='margin: 8px 0; font-size: 0.95rem; color: #1e293b; line-height: 1.5;'>{p}</p>" for p in paragraphs])

    sub_info_html = ""
    if submitted_to or parsed_sub_date:
        sub_date_formatted = parsed_sub_date.strftime('%d/%m/%Y') if parsed_sub_date else ''
        sub_info_html = f"""
        <div style="background-color: #f1f5f9; padding: 12px 16px; border-radius: 8px; margin: 15px 0; border-left: 4px solid #2563eb;">
            {'<p style="margin: 4px 0;"><strong>Submission Date:</strong> ' + sub_date_formatted + '</p>' if sub_date_formatted else ''}
            {'<p style="margin: 4px 0;"><strong>Submitted To:</strong> ' + submitted_to + '</p>' if submitted_to else ''}
        </div>
        """

    aging_legend_html = ""
    if mail_type == 'outstanding':
        aging_legend_html = """
        <div style="margin: 15px 0 12px 0; padding: 10px 14px; background-color: #f8fafc; border: 1px solid #e2e8f0; border-radius: 6px;">
            <strong style="font-size: 0.85rem; color: #334155; margin-right: 12px;">Payment Status Reference:</strong>
            <span style="display: inline-block; padding: 3px 8px; border-radius: 4px; font-size: 0.8rem; font-weight: bold; background-color: #ecfdf5; color: #059669; border: 1px solid #10b981; margin-right: 8px;">
                &#9679; Not Due (Days &le; Credit Days)
            </span>
            <span style="display: inline-block; padding: 3px 8px; border-radius: 4px; font-size: 0.8rem; font-weight: bold; background-color: #fffbeb; color: #d97706; border: 1px solid #f59e0b; margin-right: 8px;">
                &#9679; Due (&gt; Credit Days &le; 2&times; Credit Days)
            </span>
            <span style="display: inline-block; padding: 3px 8px; border-radius: 4px; font-size: 0.8rem; font-weight: bold; background-color: #fef2f2; color: #dc2626; border: 1px solid #ef4444;">
                &#9679; Overdue (&gt; 2&times; Credit Days)
            </span>
        </div>
        """

    table_rows_html = ""
    for row in preview_rows:
        inv_no_display = row['inv_no']
        status_col_td = ""
        if mail_type == 'outstanding':
            if row.get('aging'):
                ag = row['aging']
                inv_no_display = f"""<span style="display: inline-block; padding: 3px 8px; border-radius: 4px; font-weight: bold; font-size: 0.85rem; background-color: {ag['bg_color']}; color: {ag['color']}; border: 1px solid {ag['border_color']};">{row['inv_no']}</span>"""
                status_col_td = f"""<td style="padding: 8px 12px; border: 1px solid #cbd5e1; text-align: center;"><span style="color: {ag['color']}; font-weight: bold; font-size: 0.85rem;">{ag['status_label']}</span></td>"""
            else:
                status_col_td = """<td style="padding: 8px 12px; border: 1px solid #cbd5e1; text-align: center;">-</td>"""

        table_rows_html += f"""
        <tr>
            <td style="padding: 8px 12px; border: 1px solid #cbd5e1;">{inv_no_display}</td>
            <td style="padding: 8px 12px; border: 1px solid #cbd5e1;">{row['inv_date']}</td>
            <td style="padding: 8px 12px; border: 1px solid #cbd5e1;">{row['customer']}</td>
            <td style="padding: 8px 12px; border: 1px solid #cbd5e1;">{row['department']}</td>
            {status_col_td}
            <td style="padding: 8px 12px; border: 1px solid #cbd5e1; text-align: right; font-weight: bold;">&#8377; {row['amount']:,.2f}</td>
        </tr>
        """

    status_header_th = """<th style="padding: 8px 12px; border: 1px solid #001f3f; text-align: center;">Status</th>""" if mail_type == 'outstanding' else ""
    colspan_val = 5 if mail_type == 'outstanding' else 4

    html_body = f"""
    <html>
    <body style="font-family: Arial, sans-serif; color: #1e293b; line-height: 1.5; padding: 10px;">
        <h2 style="color: #001f3f; margin-bottom: 16px;">{subject}</h2>
        {formatted_body_html}
        {sub_info_html}
        {aging_legend_html}
        <h4 style="color: #0f172a; margin-top: 15px; margin-bottom: 8px;">Invoice Summary:</h4>
        <table style="width: 100%; border-collapse: collapse; margin-top: 5px; font-size: 0.9rem;">
            <thead>
                <tr style="background-color: #001f3f; color: #ffffff;">
                    <th style="padding: 8px 12px; border: 1px solid #001f3f; text-align: left;">Invoice No</th>
                    <th style="padding: 8px 12px; border: 1px solid #001f3f; text-align: left;">Invoice Date</th>
                    <th style="padding: 8px 12px; border: 1px solid #001f3f; text-align: left;">Customer</th>
                    <th style="padding: 8px 12px; border: 1px solid #001f3f; text-align: left;">Department</th>
                    {status_header_th}
                    <th style="padding: 8px 12px; border: 1px solid #001f3f; text-align: right;">Amount (INR)</th>
                </tr>
            </thead>
            <tbody>
                {table_rows_html}
            </tbody>
            <tfoot>
                <tr style="background-color: #f8fafc; font-weight: bold;">
                    <td colspan="{colspan_val}" style="padding: 8px 12px; border: 1px solid #cbd5e1; text-align: right;">Total Amount:</td>
                    <td style="padding: 8px 12px; border: 1px solid #cbd5e1; text-align: right; color: #059669;">&#8377; {total_amt:,.2f}</td>
                </tr>
            </tfoot>
        </table>
        <br>
        <p style="font-size: 0.85rem; color: #64748b;">This is an automated notification from BVM Management System.</p>
    </body>
    </html>
    """

    email_sent = False
    try:
        from django.core.mail import EmailMessage
        from django.conf import settings

        from_email = getattr(settings, 'DEFAULT_FROM_EMAIL', None) or getattr(settings, 'EMAIL_HOST_USER', 'noreply@bvm.com')
        msg = EmailMessage(
            subject=subject,
            body=html_body,
            from_email=from_email,
            to=to_recipients,
            cc=cc_recipients
        )
        msg.content_subtype = 'html'

        # Attach combined PDF ONLY for invoice submission (NOT for outstanding followup)
        if mail_type != 'outstanding' and trans_inv_nos:
            try:
                pdf_bytes = _build_combined_pdf_for_invoices(trans_inv_nos)
                if pdf_bytes:
                    clean_first = trans_inv_nos[0].replace('/', '_').replace('\\', '_')
                    pdf_filename = "Combined_Invoices.pdf" if len(trans_inv_nos) > 1 else f"Invoice_{clean_first}.pdf"
                    msg.attach(pdf_filename, pdf_bytes, 'application/pdf')
            except Exception as e:
                print('Error attaching combined PDF:', e)

        msg.send(fail_silently=True)
        email_sent = True
    except Exception as e:
        print('Error sending invoice submission email:', e)

    if mail_type == 'outstanding':
        msg_str = f'Successfully sent Outstanding Followup email for {len(preview_rows)} invoice(s).'
    else:
        msg_str = f'Successfully updated submission details for {updated_count} invoice(s).'
    if email_sent:
        msg_str += ' Email sent to recipients.'

    return JsonResponse({
        'success': True,
        'msg': msg_str,
        'mail_type': mail_type,
        'submission_date': parsed_sub_date.strftime('%Y-%m-%d') if parsed_sub_date else '',
        'submitted_to': submitted_to
    })


@login_required(login_url='login_page')
def ar_preview_combined_pdf(request):
    """
    Renders and streams inline the combined PDF for the selected invoice IDs.
    """
    raw_ids = request.GET.get('ids', '')
    if not raw_ids:
        return HttpResponse("No invoice IDs provided.", content_type="text/plain")

    id_list = [x.strip() for x in raw_ids.split(',') if x.strip()]
    inv_nos = []
    for item_id in id_list:
        if item_id.startswith('trans_'):
            try:
                ti_id = int(item_id.replace('trans_', ''))
                ti = TransInvoiceInfo.objects.filter(id=ti_id).first()
                if ti and ti.ti_inv_no and ti.ti_inv_no not in inv_nos:
                    inv_nos.append(ti.ti_inv_no)
            except Exception:
                pass
        else:
            if item_id not in inv_nos:
                inv_nos.append(item_id)

    if not inv_nos:
        return HttpResponse("No valid invoices found to preview.", content_type="text/plain")

    pdf_bytes = _build_combined_pdf_for_invoices(inv_nos)
    if not pdf_bytes:
        return HttpResponse("No combined PDF documents available for the selected invoice(s).", content_type="text/plain")

    clean_first = inv_nos[0].replace('/', '_').replace('\\', '_')
    filename = "Combined_Invoices_Preview.pdf" if len(inv_nos) > 1 else f"Invoice_{clean_first}.pdf"
    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}"'
    return response


@login_required(login_url='login_page')
def ar_export_selected_excel(request):
    """
    Export selected invoices to Excel with submission date + recipient.
    """
    if request.method != "POST":
        return redirect('ar_list')

    raw_ids = request.POST.get('selected_ids', '')
    submitted_date_str = request.POST.get('submitted_date', '').strip()
    recipient = request.POST.get('recipient', '').strip()

    id_list = [x.strip() for x in raw_ids.split(',') if x.strip()]
    if not id_list:
        return redirect('ar_list')

    parsed_sub_date = None
    if submitted_date_str:
        try:
            parsed_sub_date = datetime.strptime(submitted_date_str, '%Y-%m-%d').date()
        except Exception:
            parsed_sub_date = timezone.localtime(timezone.now()).date()

    rows_to_export = []
    today = timezone.localtime(timezone.now()).date()

    for item_id in id_list:
        if item_id.startswith('trans_'):
            try:
                ti_id = int(item_id.replace('trans_', ''))
                ti = TransInvoiceInfo.objects.select_related(
                    'ti_customer', 'ti_customer__cu_business_sol',
                    'ti_trip', 'ti_trip__tr_enquirynumber',
                    'ti_consignment', 'ti_consignment__co_enquirynumber',
                ).get(id=ti_id)

                op_date = None
                try:
                    # 1. From saved line items (most recent planning date)
                    woh_dts = (
                        TransInvoiceInfo.objects
                        .filter(is_woh=True, ti_inv_no=ti.ti_inv_no)
                        .values_list(
                            'ti_trip__tr_enquirynumber__en_pickupdatetime',
                            'ti_consignment__co_enquirynumber__en_pickupdatetime'
                        )
                    )
                    valid_dates = []
                    for trip_dt, cons_dt in woh_dts:
                        dt = trip_dt or cons_dt
                        if dt:
                            valid_dates.append(dt.date() if hasattr(dt, 'date') else dt)

                    if valid_dates:
                        op_date = max(valid_dates)
                    elif ti.ti_trip and ti.ti_trip.tr_enquirynumber and ti.ti_trip.tr_enquirynumber.en_pickupdatetime:
                        pickup_dt = ti.ti_trip.tr_enquirynumber.en_pickupdatetime
                        op_date = pickup_dt.date() if hasattr(pickup_dt, 'date') else pickup_dt
                    elif ti.ti_consignment and ti.ti_consignment.co_enquirynumber and ti.ti_consignment.co_enquirynumber.en_pickupdatetime:
                        pickup_dt = ti.ti_consignment.co_enquirynumber.en_pickupdatetime
                        op_date = pickup_dt.date() if hasattr(pickup_dt, 'date') else pickup_dt
                except Exception:
                    op_date = None

                if parsed_sub_date:
                    try:
                        with connection.cursor() as cursor:
                            cursor.execute(
                                "UPDATE sms_app_transinvoiceinfo SET ti_submission_date = %s, ti_submitted_to = %s WHERE ti_inv_no = %s",
                                [parsed_sub_date, recipient, ti.ti_inv_no]
                            )
                    except Exception as e:
                        print('Error updating trans invoice submission via SQL:', e)

                comp_name = 'BVM Trans Solutions Pvt Ltd'
                if ti.ti_customer and ti.ti_customer.cu_business_sol:
                    comp_name = str(ti.ti_customer.cu_business_sol)

                c_name = (
                    ti.ti_customer_short_name or
                    (getattr(ti.ti_customer, 'cu_nameshort', None) if ti.ti_customer else '') or
                    (str(ti.ti_customer) if ti.ti_customer else '')
                )
                tot = float(ti.ti_total or 0.0)
                svc = tot  # Service Value = ti_total
                cgst = 0.0
                sgst = 0.0
                igst = 0.0

                due_op = (today - op_date).days if op_date else ''
                due_sub = ''  # Kept empty per requirement

                # Customer department from WOH line items
                c_dept = ti.ti_department or ''
                if not c_dept:
                    woh_dept = (
                        TransInvoiceInfo.objects
                        .filter(is_woh=True, ti_inv_no=ti.ti_inv_no)
                        .exclude(ti_department__isnull=True)
                        .exclude(ti_department='')
                        .values_list('ti_department', flat=True)
                        .first()
                    )
                    c_dept = woh_dept or ''

                rows_to_export.append({
                    'company': comp_name,
                    'product': 'Transport',
                    'branch': ti.ti_branch or '',
                    'operation_date': op_date.strftime('%d-%m-%Y') if op_date else '',
                    'invoice_no': ti.ti_inv_no or '',
                    'invoice_date': ti.ti_inv_date.strftime('%d-%m-%Y') if ti.ti_inv_date else '',
                    'customer': c_name,
                    'department': c_dept,
                    'service_value': svc,
                    'cgst': cgst,
                    'sgst': sgst,
                    'igst': igst,
                    'total_amount': svc + cgst + sgst + igst,
                    'due_op': due_op,
                    'due_sub': due_sub,
                    'due_payment_received_date': '',
                    'payment_received_amount': 0.0,
                    'tds': 0.0,
                    'sales_person': '',
                    'status': 'Transport Invoice',
                })
            except Exception as e:
                print('Error processing trans:', e)

        elif item_id.startswith('ar_'):
            try:
                ar_id = int(item_id.replace('ar_', ''))
                ar = Ar_Info.objects.select_related(
                    'ar_company', 'ar_product', 'ar_branch', 'ar_invoice_num',
                    'ar_customer_name', 'ar_customer_dept', 'ar_status', 'ar_sales_person'
                ).get(id=ar_id)

                if parsed_sub_date:
                    ar.ar_submission_date = parsed_sub_date
                    if ar.ar_operation_date:
                        ar.ar_due_from_operation_date = (parsed_sub_date - ar.ar_operation_date).days
                if recipient:
                    ar.ar_invoice_sent_to = recipient
                ar.save()

                tot = float(ar.ar_amount or 0.0)
                svc = float(ar.ar_service_value or 0.0)

                rows_to_export.append({
                    'company': str(ar.ar_company) if ar.ar_company else 'BVM Pack Solutions pvt ltd',
                    'product': str(ar.ar_product) if ar.ar_product else 'Packing',
                    'branch': str(ar.ar_branch) if ar.ar_branch else '',
                    'operation_date': ar.ar_operation_date.strftime('%d-%m-%Y') if ar.ar_operation_date else '',
                    'invoice_no': str(ar.ar_invoice_num or ''),
                    'invoice_date': ar.ar_invoice_date.strftime('%d-%m-%Y') if ar.ar_invoice_date else '',
                    'customer': str(ar.ar_customer_name or ''),
                    'department': str(ar.ar_customer_dept or ''),
                    'service_value': svc,
                    'cgst': float(ar.ar_cgst or 0.0),
                    'sgst': float(ar.ar_sgst or 0.0),
                    'igst': float(ar.ar_igst or 0.0),
                    'total_amount': tot,
                    'due_op': ar.ar_due_from_operation_date or '',
                    'due_sub': '',
                    'due_payment_received_date': ar.ar_payment_received_date.strftime('%d-%m-%Y') if ar.ar_payment_received_date else '',
                    'payment_received_amount': float(ar.ar_payment_received_amount or 0.0),
                    'tds': float(ar.ar_tds or 0.0),
                    'sales_person': str(ar.ar_sales_person or ''),
                    'status': str(ar.ar_status) if ar.ar_status else 'Submitted',
                })
            except Exception as e:
                print('Error processing ar:', e)

    # Generate Excel Workbook matching exact user fields
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AR Invoices"
    ws.views.sheetView[0].showGridLines = True

    # Title Block
    ws.merge_cells('A1:R1')
    title_cell = ws['A1']
    title_cell.value = "ACCOUNTS RECEIVABLE - INVOICE SUBMISSION REPORT"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="001F3F", end_color="001F3F", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Subtitle / Metadata Block (Includes submitted date & recipient)
    ws.merge_cells('A2:R2')
    sub_cell = ws['A2']
    sub_cell.value = f"Export Date: {today.strftime('%d-%m-%Y')}  |  Total Invoices: {len(rows_to_export)}  |  Submission Date: {submitted_date_str or 'N/A'}  |  Recipient: {recipient or 'N/A'}"
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="4A5568")
    sub_cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    # Headers matching the exact user fields
    headers = [
        "S.No", "Company", "Product", "Branch", "Operation Date",
        "Invoice Number", "Invoice Date", "Customer Name", "Customer Department",
        "Service Value", "CGST", "SGST", "IGST", "Amount (INR)",
        "Due From Operation Date", "Due From Submission Date", "Sales Executive", "Status"
    ]

    # Header Row (Row 4)
    ws.row_dimensions[3].height = 10  # spacing
    ws.row_dimensions[4].height = 28

    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_num, h_text in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Data Rows
    current_row = 5
    tot_svc = 0.0
    tot_amt = 0.0

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    for idx, r in enumerate(rows_to_export, 1):
        row_fill = zebra_fill if idx % 2 == 0 else white_fill
        ws.row_dimensions[current_row].height = 22

        svc_val = r['service_value']
        amt_val = r['total_amount']

        tot_svc += svc_val
        tot_amt += amt_val

        vals = [
            idx,
            r['company'],
            r['product'],
            r['branch'],
            r['operation_date'],
            r['invoice_no'],
            r['invoice_date'],
            r['customer'],
            r['department'],
            svc_val,
            r['cgst'],
            r['sgst'],
            r['igst'],
            amt_val,
            r['due_op'],
            r['due_sub'],
            r['sales_person'],
            r['status'],
        ]

        for col_idx, val in enumerate(vals, 1):
            c = ws.cell(row=current_row, column=col_idx, value=val)
            c.fill = row_fill
            c.border = thin_border
            c.font = Font(name="Calibri", size=9)

            # Financial number format
            if col_idx in [10, 11, 12, 13, 14]:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = '#,##0.00'
            elif col_idx in [1, 5, 7, 15, 16, 18]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

        current_row += 1

    # Total Row
    ws.row_dimensions[current_row].height = 25
    total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    total_font = Font(name="Calibri", size=10, bold=True, color="0F172A")

    ws.cell(row=current_row, column=1, value="TOTAL")
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
    for ci in range(1, 10):
        cell = ws.cell(row=current_row, column=ci)
        cell.fill = total_fill
        cell.font = total_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    svc_cell = ws.cell(row=current_row, column=10, value=tot_svc)
    svc_cell.fill = total_fill
    svc_cell.font = total_font
    svc_cell.number_format = '#,##0.00'
    svc_cell.border = thin_border

    for ci in range(11, 14):
        cell = ws.cell(row=current_row, column=ci, value="-")
        cell.fill = total_fill
        cell.font = total_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    amt_cell = ws.cell(row=current_row, column=14, value=tot_amt)
    amt_cell.fill = total_fill
    amt_cell.font = total_font
    amt_cell.number_format = '#,##0.00'
    amt_cell.border = thin_border

    for ci in range(15, 19):
        cell = ws.cell(row=current_row, column=ci, value="")
        cell.fill = total_fill
        cell.border = thin_border

    # Adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row in [1, 2]:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    filename = f"AR_Invoices_Export_{today.strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required(login_url='login_page')
def ar_add(request, ar_id=0):
    context = {}
    first_name = request.session.get('first_name')
    user_id = request.session.get('ses_userID')
    role = None
    try:
        role = User_extInfo.objects.get(user=user_id).emp_role
    except Exception:
        pass

    if request.method == "GET":
        if ar_id == 0:
            form = ArinfoaddForm()
            context = {
                'form': form,
                'role': role,
                'first_name': first_name,
                'user_id': user_id,
            }
        else:
            arinfo = Ar_Info.objects.get(pk=ar_id)
            invoice_number = arinfo.ar_invoice_num
            arcomments_list = Ar_comments_Info.objects.filter(arc_invoice_num=invoice_number)
            form = ArinfoaddForm(instance=arinfo)
            context = {
                'form': form,
                'role': role,
                'first_name': first_name,
                'user_id': user_id,
                'arcomments_list': arcomments_list,
            }
        return render(request, "asset_mgt_app/ar_add.html", context)
    else:
        if ar_id == 0:
            form = ArinfoaddForm(request.POST)
        else:
            arinfo = Ar_Info.objects.get(pk=ar_id)
            form = ArinfoaddForm(request.POST, instance=arinfo)
        if form.is_valid():
            form.save()
        return redirect('/SMS/ar_list')


@login_required(login_url='login_page')
def ar_delete(request, ar_id):
    arinfo = Ar_Info.objects.get(pk=ar_id)
    arinfo.delete()
    return redirect('/SMS/ar_list')